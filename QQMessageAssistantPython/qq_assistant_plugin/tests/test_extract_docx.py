from docx import Document

from core.extractors.docx import extract_docx


def test_extract_paragraphs_and_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("入团申请书")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    path = tmp_path / "t.docx"; doc.save(path)

    data = extract_docx(path)
    assert data["kind"] == "docx"
    assert data["paragraphs"][0]["text"] == "入团申请书"
    assert data["tables"][0]["rows"][0][0] == "姓名"
    assert data["tables"][0]["rows"][0][1] == ""
