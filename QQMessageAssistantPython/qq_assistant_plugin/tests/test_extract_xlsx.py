from openpyxl import Workbook

from core.extractors.xlsx import extract_xlsx


def test_extract_cells_and_merges(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名"; ws["B1"] = ""
    ws["A2"] = 100
    ws.merge_cells("A3:B3")
    path = tmp_path / "t.xlsx"; wb.save(path)

    data = extract_xlsx(path)
    assert data["kind"] == "xlsx"
    cells = {c["ref"]: c["value"] for c in data["sheets"][0]["cells"]}
    # 空单元格不输出；数字转字符串
    assert cells == {"A1": "姓名", "A2": "100"}
    assert data["sheets"][0]["merged"] == ["A3:B3"]
    assert data["sheets"][0]["maxRow"] == 3


def test_extract_date_isoformat(tmp_path):
    import datetime
    wb = Workbook()
    wb.active["A1"] = datetime.date(2026, 10, 22)
    path = tmp_path / "d.xlsx"; wb.save(path)

    data = extract_xlsx(path)
    # openpyxl 读回 date 会变 datetime，isoformat 带时间部分；Dify 能理解，不算 bug
    assert data["sheets"][0]["cells"][0]["value"].startswith("2026-10-22")
