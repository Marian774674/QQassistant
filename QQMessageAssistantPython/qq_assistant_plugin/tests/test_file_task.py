"""文件任务测试：下载/抽取与回写分离，语义决策已全部交给 Chatflow。

钉三件事：
1. ingest 只做 PENDING→PROCESSING，产出带 structure 的任务上下文
2. apply_fill 只按 Dify 的 fills 机械回写，高置信回传 / 低置信进草稿
3. 下载失败、空 fills、不可写格式都有明确兜底，不发空文件
"""
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
import pytest
import respx
from langbot_plugin.api.entities.builtin.platform import message as pm
from openpyxl import Workbook, load_workbook

from core import file_store
from core.config import Settings
from core.file_task import apply_fill, ingest

FILE_URL = "http://files.test/in.xlsx"


class FakePlugin:
    """插件替身：记录 send_message，属性对齐真 deps。"""

    def __init__(self, settings):
        self.sent = []
        self.logs = []
        self.settings = settings
        self.spring = AsyncMock()
        self.spring.profile = AsyncMock(return_value={"facts": {"姓名": "张三"}})
        self.log_queue = SimpleNamespace(enqueue=self.logs.append)

    async def send_message(self, bot_uuid, target_type, target_id, chain):
        self.sent.append((target_type, target_id, chain))

    async def call_platform_api(self, bot_uuid, action, params=None):
        return {}


class FakeCtx:
    def __init__(self, chain):
        self.event = SimpleNamespace(sender_id="1001", launcher_id=None,
                                     message_chain=chain)

    async def get_bot_uuid(self):
        return "bot-u"


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setattr(file_store, "BASE", tmp_path / "files")
    return Settings(
        spring_base_url="http://sb", spring_timeout=5,
        dify_base_url="http://dify/v1", dify_api_key="k", dify_timeout=10,
        dify_file_workflow_key="", file_confirm_threshold=0.85,
        bot_uuid="u", bot_qq=12345, group_prefix="!",
        fact_auto_threshold=0.85,
        delay_base_ms=0, delay_per_char_ms=0, delay_jitter_ms=0, delay_max_ms=0,
        intercept_refresh_s=60, intercept_scope="group",
    )


@pytest.fixture
def plugin(settings):
    return FakePlugin(settings)


def make_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "姓名"
    src = tmp_path / "src.xlsx"
    wb.save(src)
    return src.read_bytes()


def file_chain():
    return pm.MessageChain([
        pm.Plain(text="帮我把表填了"),
        pm.File(type="File", id="f1", name="名单.xlsx", size=100, url=FILE_URL),
    ])


def fill_payload(**overrides):
    payload = {"task_ref": "", "fills": [{"sheet": "Sheet1", "ref": "B1", "value": "张三"}],
               "missing": [], "confidence": 0.95, "needs_confirm": False}
    payload.update(overrides)
    return payload


@pytest.mark.asyncio
@respx.mock
async def test_ingest_extracts_structure(plugin, tmp_path):
    respx.get(FILE_URL).mock(return_value=httpx.Response(200, content=make_xlsx(tmp_path)))
    ctx = FakeCtx(file_chain())

    with patch("core.file_task.reply_text", new=AsyncMock()):
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert task is not None
    assert task["status"] == "PROCESSING"
    assert task["originalName"] == "名单.xlsx"
    assert task["structure"]["kind"] == "xlsx"
    # 任务落盘，文本补充时可直接 load_task
    assert (file_store.task_dir(task["taskRef"]) / "task.json").exists()
    statuses = [c.args[0]["status"] for c in plugin.spring.sync_file_task.await_args_list]
    assert statuses == ["PENDING", "PROCESSING"]
    # 下载阶段不再擅自回复“收到”，话术由 Chatflow 统一决定
    assert not plugin.sent


@pytest.mark.asyncio
async def test_apply_fill_writes_and_sends(plugin, tmp_path, settings):
    # 造一个已就绪的任务
    src = tmp_path / "files" / "ft_1" / "input.xlsx"
    src.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.save(src)
    task = {"taskRef": "ft_1", "senderId": 1001, "chatId": None,
            "status": "PROCESSING", "input": {"path": str(src)},
            "originalName": "名单.xlsx", "structure": {"kind": "xlsx"}}
    ctx = FakeCtx(pm.MessageChain([]))
    payload = fill_payload(task_ref="ft_1",
                           fills='[{"sheet":"Sheet1","ref":"B1","value":"张三"}]')

    with patch("core.file_task.reply_text", new=AsyncMock()):
        result = await apply_fill(task, payload, plugin, ctx)

    assert result == "SENT"
    out = file_store.task_dir("ft_1") / "名单_filled.xlsx"
    assert out.exists()
    assert load_workbook(out)["Sheet1"]["B1"].value == "张三"
    assert plugin.sent
    plugin.spring.create_draft.assert_not_awaited()
    # 文件是插件直接发的，必须自己补 OUT 日志（否则 message_log 只有文本）
    assert plugin.logs[-1]["msgType"] == "file"
    assert plugin.logs[-1]["direction"] == "OUT"
    assert plugin.logs[-1]["content"] == "名单_filled.xlsx"


@pytest.mark.asyncio
async def test_apply_fill_low_confidence_creates_draft(plugin, tmp_path):
    src = tmp_path / "files" / "ft_2" / "input.xlsx"
    src.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.save(src)
    task = {"taskRef": "ft_2", "senderId": 1001, "chatId": None,
            "status": "PROCESSING", "input": {"path": str(src)},
            "originalName": "名单.xlsx", "structure": {"kind": "xlsx"}}
    ctx = FakeCtx(pm.MessageChain([]))

    with patch("core.file_task.reply_text", new=AsyncMock()):
        result = await apply_fill(task, fill_payload(confidence=0.4), plugin, ctx)

    assert result == "AWAITING_CONFIRM"
    plugin.spring.create_draft.assert_awaited_once()
    payload = plugin.spring.create_draft.await_args.args[4]
    assert payload["taskRef"] == "ft_2"
    assert payload["fileRef"].startswith("local://")
    # 草稿必须自带目标：否则 Spring 会在群/私聊之间猜错
    assert payload["targetType"] == "private"
    assert payload["targetId"] == 1001
    assert not plugin.sent


@pytest.mark.asyncio
async def test_group_task_draft_targets_group(plugin, tmp_path):
    src = tmp_path / "files" / "ft_g" / "input.xlsx"
    src.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.save(src)
    task = {"taskRef": "ft_g", "senderId": 1001, "chatId": 2002,
            "status": "PROCESSING", "input": {"path": str(src)},
            "originalName": "名单.xlsx", "structure": {"kind": "xlsx"}}
    ctx = FakeCtx(pm.MessageChain([]))

    with patch("core.file_task.reply_text", new=AsyncMock()):
        await apply_fill(task, fill_payload(confidence=0.4), plugin, ctx)

    payload = plugin.spring.create_draft.await_args.args[4]
    assert payload["targetType"] == "group"
    assert payload["targetId"] == 2002


@pytest.mark.asyncio
async def test_sync_carries_phase3_fields(plugin, tmp_path):
    """Spring 的 file_task 要能看到会话、要求、缺口、产物引用。"""
    task = {
        "taskRef": "ft_9", "senderId": 1001, "chatId": None,
        "status": "AWAITING_CONFIRM",
        "input": {"path": str(tmp_path / "in.xlsx")},
        "output": {"path": str(tmp_path / "out.xlsx")},
        "request": "把姓名填上", "conversationId": "conv-9",
        "missing": ["学号"], "confidence": 0.6, "needs_confirm": True,
    }

    from core.file_task import sync

    await sync(plugin, task)

    body = plugin.spring.sync_file_task.await_args.args[0]
    assert body["conversationId"] == "conv-9"
    assert body["request"] == "把姓名填上"
    assert body["missing"] == ["学号"]
    assert body["confidence"] == 0.6
    assert body["needsConfirm"] is True
    assert body["fileRef"].startswith("local://")
    assert body["outputRef"].startswith("local://")


@pytest.mark.asyncio
async def test_apply_fill_string_false_is_false(plugin, tmp_path):
    """Dify 把布尔值输出成字符串 "false" 时，不能 bool("false") 误判成需要确认。"""
    src = tmp_path / "files" / "ft_3" / "input.xlsx"
    src.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.save(src)
    task = {"taskRef": "ft_3", "senderId": 1001, "chatId": None,
            "status": "PROCESSING", "input": {"path": str(src)},
            "originalName": "名单.xlsx", "structure": {"kind": "xlsx"}}
    ctx = FakeCtx(pm.MessageChain([]))

    with patch("core.file_task.reply_text", new=AsyncMock()):
        result = await apply_fill(
            task, fill_payload(needs_confirm="false"), plugin, ctx)

    assert result == "SENT"
    plugin.spring.create_draft.assert_not_awaited()


@pytest.mark.asyncio
async def test_apply_fill_empty_fills_reports_structure(plugin, tmp_path):
    src = tmp_path / "files" / "ft_4" / "input.xlsx"
    src.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.save(src)
    task = {"taskRef": "ft_4", "senderId": 1001, "chatId": None,
            "status": "PROCESSING", "input": {"path": str(src)},
            "originalName": "名单.xlsx",
            "structure": {"kind": "xlsx", "sheets": [
                {"name": "Sheet1", "cells": [{"ref": "A1", "value": "姓名"}]}]}}
    ctx = FakeCtx(pm.MessageChain([]))

    with patch("core.file_task.reply_text", new=AsyncMock()) as rt:
        result = await apply_fill(task, fill_payload(fills=[]), plugin, ctx)

    assert result == "FAILED"
    assert any("没找到要填的内容" in c.args[1] and "A1=姓名" in c.args[1]
               for c in rt.await_args_list)
    assert not plugin.sent


@pytest.mark.asyncio
@respx.mock
async def test_unwritable_format_fails_fast(plugin):
    pdf_chain = pm.MessageChain([
        pm.File(type="File", id="f2", name="材料.pdf", size=10, url=FILE_URL)])
    respx.get(FILE_URL).mock(return_value=httpx.Response(200, content=b"%PDF-fake"))
    ctx = FakeCtx(pdf_chain)

    with patch("core.file_task.reply_text", new=AsyncMock()) as rt:
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert task is None
    assert plugin.spring.sync_file_task.await_args_list[-1].args[0]["status"] == "FAILED"
    assert any("只能读不能改" in c.args[1] for c in rt.await_args_list)


@pytest.mark.asyncio
@respx.mock
async def test_download_via_napcat_http(settings, tmp_path):
    """NapCat 回传的 File 没有 url（实机 0.2.4 现象），必须走 NapCat HTTP API 换址。"""
    s = replace(settings, napcat_api_url="http://napcat:3000", napcat_token="tk")
    plugin = FakePlugin(s)
    respx.post("http://napcat:3000/get_private_file_url").mock(
        return_value=httpx.Response(200, json={"status": "ok", "data": {"url": FILE_URL}}))
    respx.get(FILE_URL).mock(return_value=httpx.Response(200, content=make_xlsx(tmp_path)))
    chain = pm.MessageChain([
        pm.Plain(text="帮我填"),
        pm.File(type="File", id="fid_1", name="名单.xlsx", size=1, url=""),
    ])
    ctx = FakeCtx(chain)

    with patch("core.file_task.reply_text", new=AsyncMock()):
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert task is not None
    assert task["structure"]["kind"] == "xlsx"


@pytest.mark.asyncio
@respx.mock
async def test_download_502_refetches_url_and_succeeds(settings, tmp_path):
    """QQ CDN 某个边缘节点 502 时，换 NapCat 新 URL 重试成功。"""
    retry_url = "http://files.test/retry.xlsx"
    s = replace(settings, napcat_api_url="http://napcat:3000", napcat_token="")
    plugin = FakePlugin(s)
    respx.get(FILE_URL).mock(return_value=httpx.Response(502, text="bad gateway"))
    respx.post("http://napcat:3000/get_private_file_url").mock(
        return_value=httpx.Response(200, json={"status": "ok", "data": {"url": retry_url}}))
    respx.get(retry_url).mock(return_value=httpx.Response(200, content=make_xlsx(tmp_path)))
    ctx = FakeCtx(file_chain())

    with patch("core.file_task.reply_text", new=AsyncMock()):
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert respx.get(FILE_URL).called
    assert respx.get(retry_url).called
    assert task is not None


@pytest.mark.asyncio
@respx.mock
async def test_download_502_exhausts_three_attempts(settings, tmp_path):
    """连续 502 时最多试 3 次，失败文案要带状态码，且不能把签名 URL 全量回给用户。"""
    retry_url = "http://files.test/retry.xlsx"
    s = replace(settings, napcat_api_url="http://napcat:3000", napcat_token="")
    plugin = FakePlugin(s)
    respx.get(FILE_URL).mock(return_value=httpx.Response(502, text="bad gateway"))
    respx.post("http://napcat:3000/get_private_file_url").mock(
        return_value=httpx.Response(200, json={"status": "ok", "data": {"url": retry_url}}))
    respx.get(retry_url).mock(return_value=httpx.Response(502, text="bad gateway"))
    ctx = FakeCtx(file_chain())

    with patch("core.file_task.reply_text", new=AsyncMock()) as rt:
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert task is None
    synced = plugin.spring.sync_file_task.await_args_list[-1].args[0]
    assert synced["status"] == "FAILED"
    assert "502" in synced["detail"]["reason"]
    assert any("502" in c.args[1] for c in rt.await_args_list)
    assert all("rkey" not in c.args[1] for c in rt.await_args_list)


@pytest.mark.asyncio
@respx.mock
async def test_napcat_502_reports_real_reason(settings):
    """NapCat 502（服务没起/被系统代理接管）时，失败文案要带真实原因。"""
    s = replace(settings, napcat_api_url="http://napcat:3000", napcat_token="")
    plugin = FakePlugin(s)
    respx.post("http://napcat:3000/get_private_file_url").mock(
        return_value=httpx.Response(502, text=""))
    chain = pm.MessageChain([
        pm.File(type="File", id="fid_9", name="名单.xlsx", size=1, url="")])
    ctx = FakeCtx(chain)

    with patch("core.file_task.reply_text", new=AsyncMock()) as rt:
        task = await ingest(ctx, ctx.event, is_group=False, deps=plugin)

    assert task is None
    synced = plugin.spring.sync_file_task.await_args_list[-1].args[0]
    assert synced["status"] == "FAILED"
    assert "502" in synced["detail"]["reason"] and "网关" in synced["detail"]["reason"]
    assert any("502" in c.args[1] for c in rt.await_args_list)
