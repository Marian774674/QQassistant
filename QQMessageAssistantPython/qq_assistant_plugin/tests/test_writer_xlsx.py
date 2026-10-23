from openpyxl import Workbook, load_workbook

from core.writers.xlsx import apply_xlsx


def test_apply_fills_by_ref(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    src = tmp_path / "in.xlsx"; wb.save(src)
    out = tmp_path / "out.xlsx"

    errors = apply_xlsx(src, out, [
        {"sheet": "Sheet1", "ref": "B2", "value": "张三"},
        {"sheet": "Nope", "ref": "A1", "value": "x"},     # 忽略并记错误
    ])

    assert load_workbook(out).active["B2"].value == "张三"
    assert len(errors) == 1
    # 原文件不被覆盖（全局约束）
    assert load_workbook(src).active["B2"].value is None
