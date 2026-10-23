from docx import Document


def extract_docx(path) -> dict:
    doc = Document(path)
    paragraphs = [{"index": i, "text": p.text} for i, p in enumerate(doc.paragraphs)]
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append({"index": ti, "rows": rows})
    return {"kind": "docx", "paragraphs": paragraphs, "tables": tables}
