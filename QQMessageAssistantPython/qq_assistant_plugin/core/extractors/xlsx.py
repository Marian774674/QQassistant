from openpyxl import load_workbook


def extract_xlsx(path) -> dict:
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = cell.value.isoformat() if hasattr(cell.value, "isoformat") else str(cell.value)
                if text.strip():
                    cells.append({"ref": cell.coordinate, "value": text})
        sheets.append({
            "name": ws.title,
            "maxRow": ws.max_row,
            "maxCol": ws.max_column,
            "cells": cells,
            "merged": [str(r) for r in ws.merged_cells.ranges],
        })
    return {"kind": "xlsx", "sheets": sheets}
