from docx import Document


def apply_docx(src, out, fills) -> list[str]:
    doc = Document(src)
    errors = []
    for fill in fills or []:
        kind = fill.get("kind")
        try:
            if kind == "paragraph":
                doc.paragraphs[int(fill["index"])].text = str(fill.get("value", ""))
            elif kind == "table":
                table = doc.tables[int(fill["table"])]
                table.cell(int(fill["row"]), int(fill["col"])).text = str(fill.get("value", ""))
            else:
                errors.append(f"unknown kind: {kind}")
        except (IndexError, KeyError, ValueError) as e:
            errors.append(f"bad fill {fill}: {e}")
    doc.save(out)
    return errors
