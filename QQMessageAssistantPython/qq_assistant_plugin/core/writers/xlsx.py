from openpyxl import load_workbook


def apply_xlsx(src, out, fills) -> list[str]:
    wb = load_workbook(src)
    errors = []
    for fill in fills or []:
        name = fill.get("sheet")
        ref = fill.get("ref")
        if not name or not ref:
            errors.append(f"missing sheet/ref: {fill}")
            continue
        if name not in wb.sheetnames:
            errors.append(f"sheet not found: {name}")
            continue
        try:
            wb[name][ref] = fill.get("value")
        except Exception as e:
            # 一个填错坐标不应让整份文件处理失败；记录后继续填其它格子
            errors.append(f"bad ref {name}!{ref}: {e}")
    wb.save(out)
    return errors
